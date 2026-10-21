#!/usr/bin/env python3
"""진행자 전용 정답 구현. 참가자에게 보여 주지 않는다. 실행 위치: seed/monthly-report/"""
import shutil, sys
sys.path.insert(0, "scripts")
from openpyxl import load_workbook
from reconcile import expected, TEMPLATE
shutil.copy(TEMPLATE, "output/월간보고.xlsx")
wb = load_workbook("output/월간보고.xlsx"); ws = wb["월간보고"]
for i, row in enumerate(expected()):
    for j, v in enumerate(row[1:]):
        ws.cell(row=4 + i, column=2 + j, value=v)
ws["A2"] = "점검월: " + (sys.argv[1] if len(sys.argv) > 1 else "2026-09") + " / 노란 셀만 입력하세요"
wb.save("output/월간보고.xlsx"); print("정답 보고서 생성")
