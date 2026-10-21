#!/usr/bin/env python3
"""가상 데이터와 양식을 만든다. 모든 값은 연습용이며 실제 고객·내부 자료가 아니다."""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BRANCHES = {
    "gangnam": ("강남지점", [("예금", 412, 3820, 3, 2), ("적금", 275, 1240, 1, 0), ("펀드", 96, 2310, 7, 4),
                            ("보험", 188, 1675, 9, 5), ("대출", 134, 5210, 2, 3), ("카드", 521, 980, 4, 6)]),
    "busan": ("부산지점", [("예금", 366, 3105, 2, 1), ("적금", 241, 1090, 0, 1), ("펀드", 71, 1580, 5, 3),
                           ("보험", 152, 1320, 6, 4), ("대출", 118, 4470, 1, 2), ("카드", 433, 760, 3, 5)]),
    "daejeon": ("대전지점", [("예금", 298, 2480, 1, 1), ("적금", 205, 910, 1, 0), ("펀드", 54, 1130, 4, 2),
                             ("보험", 129, 1110, 5, 3), ("대출", 97, 3560, 1, 1), ("카드", 377, 640, 2, 4)]),
}
PREV = [("강남지점", 1630, 24, 19), ("부산지점", 1300, 15, 14), ("대전지점", 1105, 11, 9)]
NEXT = "--next" in sys.argv  # 다음 달 데이터: 이번 달을 전월로 삼고 숫자를 바꾼다(재사용 연습용)
if NEXT:
    PREV = [("강남지점", 1626, 26, 20), ("부산지점", 1381, 17, 16), ("대전지점", 1160, 14, 11)]
    for k, (nm, rows) in list(BRANCHES.items()):
        BRANCHES[k] = (nm, [(a, b + 7 * (i + 1), c + 90, d + (i % 2), e + ((i + 1) % 2)) for i, (a, b, c, d, e) in enumerate(rows)])
MONTH = "2026-10" if NEXT else "2026-09"
HEAD = ["상품", "판매건수", "판매금액(백만원)", "의심건수", "민원건수"]

os.makedirs("data", exist_ok=True)
os.makedirs("template", exist_ok=True)
os.makedirs("output", exist_ok=True)
for key, (_, rows) in BRANCHES.items():
    with open(f"data/{key}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(HEAD); w.writerows(rows)
with open("data/prev_month.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["지점", "판매건수", "의심건수", "민원건수"]); w.writerows(PREV)

wb = Workbook(); ws = wb.active; ws.title = "월간보고"
thin = Side(style="thin", color="999999"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.merge_cells("A1:G1"); ws["A1"] = "월간 소비자보호 점검 보고서 (가상)"
ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = Alignment(horizontal="center")
ws.merge_cells("A2:G2"); ws["A2"] = "점검월: 2026-09 / 노란 셀만 입력하세요"  # 양식은 항상 이 문구로 저장한다
ws["A2"].alignment = Alignment(horizontal="center")
for i, h in enumerate(["지점", "판매건수", "판매금액(백만원)", "의심건수", "민원건수", "의심비율(%)", "전월대비 판매건수 증감률(%)"], 1):
    c = ws.cell(row=3, column=i, value=h); c.font = Font(bold=True); c.border = box
    c.fill = PatternFill("solid", fgColor="D9E1F2"); c.alignment = Alignment(horizontal="center", wrap_text=True)
yellow = PatternFill("solid", fgColor="FFF2CC")
for r, name in enumerate(["강남지점", "부산지점", "대전지점", "합계"], 4):
    a = ws.cell(row=r, column=1, value=name); a.border = box; a.font = Font(bold=(name == "합계"))
    for col in range(2, 8):
        c = ws.cell(row=r, column=col); c.fill = yellow; c.border = box
        c.number_format = "0.0" if col >= 6 else "#,##0"
ws["A9"] = "※ 숫자는 지점별 CSV 합계와 일치해야 합니다."
for col, wd in zip("ABCDEFG", [14, 12, 18, 12, 12, 14, 24]): ws.column_dimensions[col].width = wd
wb.save("template/월간점검보고서_양식.xlsx")
print("점검월:", MONTH, "/ 만든 파일:", sorted(os.listdir("data")), "template/월간점검보고서_양식.xlsx")
