#!/usr/bin/env python3
"""연습용: 보고서에 일부러 숫자 오류 1개를 심는다. 사용: python3 scripts/inject_error.py output/월간보고.xlsx"""
import sys
from openpyxl import load_workbook
p = sys.argv[1]
wb = load_workbook(p); ws = wb["월간보고"]
before = ws["D7"].value; ws["D7"] = before + 3   # 합계 의심건수를 3만큼 틀리게
wb.save(p); print(f"D7 합계 의심건수 {before} -> {before + 3} (일부러 틀림)")
