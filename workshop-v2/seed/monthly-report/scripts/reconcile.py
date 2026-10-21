#!/usr/bin/env python3
"""결정적 대조 스크립트. 산술은 모델이 아니라 이 스크립트가 한다.
사용: python3 scripts/reconcile.py output/월간보고.xlsx
종료 코드: 0 일치 / 1 불일치 또는 양식 훼손 / 2 사용 오류
"""
import csv, sys, os
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "template", "월간점검보고서_양식.xlsx")
BRANCHES = [("강남지점", "gangnam"), ("부산지점", "busan"), ("대전지점", "daejeon")]
INPUT = {(r, c) for r in range(4, 8) for c in range(2, 8)}  # B4:G7만 입력 칸


def r1(x):  # 소수 첫째 자리 사사오입
    return float(Decimal(str(x)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def load_data():
    prev = {}
    with open(os.path.join(ROOT, "data", "prev_month.csv"), encoding="utf-8") as f:
        for row in csv.DictReader(f):
            prev[row["지점"]] = int(row["판매건수"])
    out = {}
    for name, key in BRANCHES:
        s = m = x = c = 0
        with open(os.path.join(ROOT, "data", f"{key}.csv"), encoding="utf-8") as f:
            for row in csv.DictReader(f):
                s += int(row["판매건수"]); m += int(row["판매금액(백만원)"]); x += int(row["의심건수"]); c += int(row["민원건수"])
        out[name] = (s, m, x, c)
    return out, prev


def expected():
    data, prev = load_data()
    rows = []
    for name, _ in BRANCHES:
        s, m, x, c = data[name]
        rows.append((name, s, m, x, c, r1(x / s * 100), r1((s - prev[name]) / prev[name] * 100)))
    ts = sum(r[1] for r in rows); tm = sum(r[2] for r in rows); tx = sum(r[3] for r in rows); tc = sum(r[4] for r in rows)
    tp = sum(prev.values())
    rows.append(("합계", ts, tm, tx, tc, r1(tx / ts * 100), r1((ts - tp) / tp * 100)))
    return rows


def main(path):
    if not os.path.exists(path):
        print(f"파일 없음: {path}", file=sys.stderr); return 2
    rep = load_workbook(path)["월간보고"]; tpl = load_workbook(TEMPLATE)["월간보고"]
    errs = []
    # 1) 숫자 대조
    labels = ["판매건수", "판매금액(백만원)", "의심건수", "민원건수", "의심비율(%)", "전월대비 판매건수 증감률(%)"]
    for i, exp in enumerate(expected()):
        r = 4 + i
        for j, label in enumerate(labels):
            cell = rep.cell(row=r, column=2 + j); want = exp[1 + j]; got = cell.value
            if not isinstance(got, (int, float)) or isinstance(got, bool) or abs(float(got) - float(want)) > 1e-9:
                errs.append(f"숫자 불일치 {cell.coordinate} [{exp[0]} {label}]: 보고서={got!r} 기대={want}")
    # 2) 양식 보존: 입력 칸 밖의 값, 병합, 열 너비, 서식
    for row in tpl.iter_rows(min_row=1, max_row=max(tpl.max_row, rep.max_row), max_col=8):
        for c in row:
            if (c.row, c.column) in INPUT:
                continue
            g = rep.cell(row=c.row, column=c.column)
            if c.coordinate == "A2":  # 점검월 문구만 바꿀 수 있다
                if not (isinstance(g.value, str) and g.value.startswith("점검월:") and g.value.endswith("/ 노란 셀만 입력하세요")):
                    errs.append(f"양식 훼손 A2: '점검월: YYYY-MM / 노란 셀만 입력하세요' 형식이어야 함 (현재 {g.value!r})")
                continue
            if g.value != c.value:
                errs.append(f"양식 훼손 {c.coordinate}: 값 {g.value!r} (양식 {c.value!r})")
    if {str(x) for x in rep.merged_cells.ranges} != {str(x) for x in tpl.merged_cells.ranges}:
        errs.append("양식 훼손: 병합 셀이 양식과 다름")
    for col in "ABCDEFG":
        if rep.column_dimensions[col].width != tpl.column_dimensions[col].width:
            errs.append(f"양식 훼손: {col}열 너비가 양식과 다름")
    for r, c in sorted(INPUT):
        a, b = rep.cell(row=r, column=c), tpl.cell(row=r, column=c)
        if a.number_format != b.number_format or a.fill.fgColor.rgb != b.fill.fgColor.rgb:
            errs.append(f"양식 훼손 {a.coordinate}: 숫자 서식 또는 채움색이 양식과 다름")
    if errs:
        print("대조 실패 (%d건)" % len(errs)); [print(" -", e) for e in errs]; return 1
    print("대조 통과: 숫자 24칸 일치, 양식 보존"); return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]) if len(sys.argv) == 2 else (print(__doc__, file=sys.stderr) or 2))
